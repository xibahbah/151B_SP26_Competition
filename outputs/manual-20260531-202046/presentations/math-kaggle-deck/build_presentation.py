from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape
from reportlab.pdfgen import canvas


OUT_DIR = Path("/Users/keith/Downloads")
PPTX_OUT = OUT_DIR / "math_reasoning_competition_presentation.pptx"
PDF_OUT = OUT_DIR / "math_reasoning_competition_presentation.pdf"

W, H = 13.333, 7.5

BLACK = RGBColor(20, 20, 20)
GRAY = RGBColor(92, 92, 92)
LIGHT = RGBColor(238, 242, 244)
ACCENT = RGBColor(112, 135, 146)
ACCENT_DARK = RGBColor(48, 62, 70)
GREEN = RGBColor(44, 132, 82)
ORANGE = RGBColor(204, 125, 54)
RED = RGBColor(170, 70, 70)


def add_text(slide, text, x, y, w, h, size=24, color=BLACK, bold=False,
             align=None, font="Helvetica", line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.name = font
    p.font.bold = bold
    p.font.color.rgb = color
    p.line_spacing = line_spacing
    if align:
        p.alignment = align
    return box


def add_bullets(slide, items, x, y, w, h, size=18, color=BLACK, gap=0.18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Helvetica"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap * 20)
        p._p.get_or_add_pPr().set("marL", "228600")
        p._p.get_or_add_pPr().set("indent", "-228600")
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.55, 0.45, 9.0, 0.45, size=17, bold=False)
    if subtitle:
        add_text(slide, subtitle, 0.55, 0.88, 9.0, 0.33, size=9.5, color=GRAY)


def add_section_slide(slide, title):
    add_text(slide, title, 0, 3.25, W, 0.55, size=28, align=PP_ALIGN.CENTER)


def rect(slide, x, y, w, h, fill=LIGHT, line=LIGHT, radius=False):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def circle(slide, x, y, d, fill=ACCENT, line=ACCENT):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def table(slide, headers, rows, x, y, w, h, col_widths=None, font_size=11):
    t = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)).table
    if col_widths:
        for i, cw in enumerate(col_widths):
            t.columns[i].width = Inches(cw)
    for c, header in enumerate(headers):
        cell = t.cell(0, c)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_DARK
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(font_size)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(248, 250, 251) if r % 2 else RGBColor(236, 241, 243)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(font_size)
                p.font.color.rgb = BLACK
    return t


def bar_chart(slide, labels, values, x, y, w, h, maxv=100, colors_list=None):
    colors_list = colors_list or [ACCENT] * len(values)
    n = len(values)
    gap = 0.22
    bw = (w - gap * (n - 1)) / n
    for i, (label, value) in enumerate(zip(labels, values)):
        bx = x + i * (bw + gap)
        bh = h * value / maxv
        rect(slide, bx, y + h - bh, bw, bh, fill=colors_list[i], line=colors_list[i])
        add_text(slide, f"{value:.1f}%", bx - 0.05, y + h - bh - 0.34, bw + 0.1, 0.25, size=10, color=BLACK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, bx - 0.05, y + h + 0.12, bw + 0.1, 0.46, size=9.2, color=GRAY, align=PP_ALIGN.CENTER)


slides = [
    {
        "kind": "title",
        "title": "Improving Math Reasoning Under a Token Budget",
        "subtitle": "CSE 151B/251B Math Reasoning Competition - Keith Gong",
    },
    {
        "kind": "summary",
        "title": "Summary",
        "body": [
            "I built an end-to-end vLLM pipeline for 1126 math problems split between MCQ and free-form response.",
            "The winning strategy was not generic fine-tuning; it was preserving the strong MCQ baseline while repairing free-form formatting and template-level failure modes.",
            "Public FRQ accuracy improved from 56.6% raw to 82.6% after targeted repair, and Kaggle public score moved from 0.650 to 0.671.",
            "The main lesson: for small thinking models, token budget and answer extraction are first-class modeling constraints."
        ],
        "bullets": ["Team: Keith Gong", "Core approach: base Qwen3-4B-Thinking + vLLM + targeted repair", "Learned: do not trade away reasoning tokens unless the replacement is measured"],
    },
    {"kind": "keywords"},
    {"kind": "section", "title": "Introduction"},
    {"kind": "team"},
    {"kind": "section", "title": "Methodology"},
    {"kind": "data"},
    {"kind": "model"},
    {"kind": "engineering"},
    {"kind": "section", "title": "Experiments"},
    {"kind": "experiment1"},
    {"kind": "experiment2"},
    {"kind": "section", "title": "Discussion"},
    {"kind": "strengths"},
    {"kind": "learned"},
    {"kind": "future"},
]


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

for spec in slides:
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    kind = spec["kind"]

    if kind == "title":
        add_text(slide, spec["title"], 1.25, 2.82, 10.8, 0.55, size=30, align=PP_ALIGN.CENTER)
        add_text(slide, spec["subtitle"], 1.6, 3.42, 10.1, 0.35, size=16, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(slide, "Qwen3-4B-Thinking | vLLM | FRQ repair | Kaggle submission", 2.0, 4.08, 9.4, 0.28, size=11, color=ACCENT_DARK, align=PP_ALIGN.CENTER)

    elif kind == "summary":
        add_title(slide, "Summary")
        add_text(slide, "3-4 sentence summary of the presentation", 0.72, 1.25, 4.8, 0.25, size=9.5, color=GRAY)
        add_bullets(slide, spec["body"], 0.75, 1.68, 7.1, 2.15, size=14.2)
        rect(slide, 8.45, 1.25, 3.85, 3.2, fill=RGBColor(244, 247, 248), line=RGBColor(224, 230, 233), radius=True)
        add_text(slide, "Final public score", 8.75, 1.55, 3.25, 0.25, size=12, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(slide, "0.671", 8.75, 1.9, 3.25, 0.62, size=36, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, "from 0.650 baseline", 8.75, 2.55, 3.25, 0.25, size=11, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(slide, "FRQ public repair", 8.75, 3.2, 3.25, 0.25, size=12, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(slide, "82.56%", 8.75, 3.52, 3.25, 0.42, size=24, color=ACCENT_DARK, bold=True, align=PP_ALIGN.CENTER)
        add_bullets(slide, spec["bullets"], 0.75, 5.0, 10.5, 0.95, size=14)

    elif kind == "keywords":
        add_title(slide, "Key Words")
        # Template-like circle rhythm.
        circle(slide, 2.45, 2.62, 0.72, fill=BLACK, line=BLACK)
        circle(slide, 3.55, 2.08, 1.9, fill=ACCENT, line=ACCENT)
        circle(slide, 5.85, 2.63, 0.86, fill=BLACK, line=BLACK)
        circle(slide, 7.12, 2.72, 0.62, fill=ACCENT, line=ACCENT)
        circle(slide, 8.08, 2.34, 1.32, fill=ACCENT_DARK, line=ACCENT_DARK)
        circle(slide, 9.78, 2.72, 0.62, fill=ACCENT, line=ACCENT)
        add_text(slide, "Qwen3\\nThinking\\nModel", 3.76, 2.5, 1.48, 0.9, size=16, color=RGBColor(255,255,255), align=PP_ALIGN.CENTER)
        add_text(slide, "vLLM", 5.94, 2.93, 0.68, 0.25, size=10, color=RGBColor(255,255,255), align=PP_ALIGN.CENTER)
        add_text(slide, "FRQ\\nRepair", 8.32, 2.75, 0.85, 0.5, size=13, color=RGBColor(255,255,255), align=PP_ALIGN.CENTER)
        add_text(slide, "Kaggle", 2.42, 3.48, 0.88, 0.25, size=9, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(slide, "Token\\nBudget", 7.03, 3.42, 0.88, 0.42, size=9, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(slide, "Postprocess", 9.55, 3.42, 1.1, 0.25, size=9, color=GRAY, align=PP_ALIGN.CENTER)

    elif kind == "section":
        add_section_slide(slide, spec["title"])

    elif kind == "team":
        add_title(slide, "Team Introduction")
        add_text(slide, "Keith Gong", 0.75, 1.55, 4.1, 0.48, size=22, bold=True)
        add_text(slide, "Department of Data Science\\nUniversity of California, San Diego\\nx4gong@ucsd.edu", 0.76, 2.18, 4.4, 0.9, size=13, color=GRAY)
        rect(slide, 6.25, 1.25, 5.65, 3.8, fill=RGBColor(245, 248, 249), line=RGBColor(228, 234, 237), radius=True)
        add_text(slide, "Project Role", 6.6, 1.58, 4.8, 0.35, size=18, bold=True)
        add_bullets(slide, [
            "Owned inference pipeline, prompt experiments, repair functions, and Kaggle submission assembly.",
            "Ran remote GPU jobs on RunPod/DataHub with vLLM and BitsAndBytes quantization.",
            "Kept MCQ and FRQ as separate systems after evidence showed they fail differently."
        ], 6.65, 2.08, 4.9, 1.8, size=14)

    elif kind == "data":
        add_title(slide, "Data Processing")
        add_text(slide, "The split has different answer contracts, so preprocessing keeps MCQ and FRQ separate.", 0.75, 1.15, 8.5, 0.3, size=13, color=GRAY)
        table(slide, ["Split", "Rows", "Use"], [
            ["Public", "1126", "Evaluation and controlled repair analysis"],
            ["Public MCQ", "375", "Baseline prompt, letter extraction"],
            ["Public FRQ", "751", "Prompting, postprocess, template repair"],
            ["Private", "943", "Submission generation only"],
            ["Private MCQ/FRQ", "300 / 643", "Merged into Kaggle CSV"],
        ], 0.75, 1.65, 6.15, 2.55, col_widths=[1.45, 1.1, 3.6], font_size=9.5)
        rect(slide, 7.35, 1.65, 4.85, 2.55, fill=RGBColor(246, 249, 250), line=RGBColor(226, 233, 236), radius=True)
        add_text(slide, "Processing pipeline", 7.7, 1.95, 4.1, 0.25, size=15, bold=True)
        add_bullets(slide, [
            "Build type-specific chat prompt",
            "Run batched vLLM generation",
            "Extract last boxed answer",
            "Normalize free-form formatting",
            "Repair known public-safe templates",
            "Write full response trace to CSV"
        ], 7.7, 2.36, 4.1, 1.65, size=12)

    elif kind == "model":
        add_title(slide, "Deep Learning Model")
        add_text(slide, "Base model", 0.75, 1.45, 2.2, 0.3, size=15, color=GRAY)
        add_text(slide, "Qwen3-4B-Thinking-2507", 0.75, 1.8, 4.5, 0.4, size=22, bold=True)
        add_bullets(slide, [
            "Served with vLLM on one GPU",
            "BitsAndBytes weight quantization",
            "Temperature 0.6, top-p 0.95, top-k 20",
            "Long thinking traces are useful but expensive"
        ], 0.78, 2.45, 4.8, 1.45, size=14)
        rect(slide, 6.25, 1.25, 5.7, 3.4, fill=RGBColor(246, 249, 250), line=RGBColor(226, 233, 236), radius=True)
        add_text(slide, "Model strategy", 6.6, 1.56, 4.8, 0.3, size=17, bold=True)
        table(slide, ["Component", "Decision"], [
            ["MCQ", "Keep baseline prompt and full token budget"],
            ["FRQ", "Base generation plus deterministic repair"],
            ["LoRA", "Rejected after holdout degradation"],
            ["CSV", "Submit full raw trace with final boxed answer"],
        ], 6.55, 2.05, 4.95, 1.85, col_widths=[1.35, 3.6], font_size=9.5)

    elif kind == "engineering":
        add_title(slide, "Engineering Tricks")
        add_text(slide, "Most gains came from treating inference as a production system, not just a prompt.", 0.75, 1.1, 9.6, 0.3, size=13, color=GRAY)
        for i, (label, body, color) in enumerate([
            ("Remote jobs", "nohup runs, log files, PID checks, nvidia-smi monitoring", ACCENT),
            ("Stability", "JSONL outputs, resumable scripts, explicit validation before submit", GREEN),
            ("Memory", "A40 GPU, KV cache tuning, avoided high-token runs that slowed to 8+ hours", ORANGE),
        ]):
            x = 0.75 + i * 4.1
            rect(slide, x, 1.75, 3.55, 2.55, fill=RGBColor(246, 249, 250), line=RGBColor(226, 233, 236), radius=True)
            circle(slide, x + 0.26, 2.08, 0.45, fill=color, line=color)
            add_text(slide, label, x + 0.85, 2.08, 2.45, 0.3, size=16, bold=True)
            add_text(slide, body, x + 0.35, 2.75, 2.9, 0.9, size=12.5, color=GRAY)
        add_text(slide, "Rule of thumb: anything over 5 minutes ran remotely with logs; notebooks were reserved for final reproducibility.", 1.15, 5.05, 10.9, 0.4, size=14, color=ACCENT_DARK, align=PP_ALIGN.CENTER)

    elif kind == "experiment1":
        add_title(slide, "Experiment 1", "Baseline, token budget, and repair")
        add_text(slide, "Initial 100-question milestone result", 0.85, 1.45, 4.6, 0.25, size=13, color=GRAY)
        bar_chart(slide, ["MCQ\\nbaseline", "FRQ\\nbaseline", "FRQ\\nstrict"], [71.1, 54.8, 56.5], 0.95, 2.0, 4.5, 2.05, colors_list=[ACCENT_DARK, ACCENT, GREEN])
        add_text(slide, "Full public FRQ pipeline", 6.25, 1.45, 4.6, 0.25, size=13, color=GRAY)
        bar_chart(slide, ["Raw", "Baseline\\nrepair", "Final\\nrepair"], [56.6, 77.1, 82.6], 6.35, 2.0, 4.65, 2.05, colors_list=[ACCENT, ORANGE, GREEN])
        add_bullets(slide, [
            "MCQ was highly sensitive to truncation; the safest move was leaving it unchanged.",
            "FRQ improved because many failures were extraction, expression, or known-template issues.",
            "Final public FRQ repair: 620/751 = 82.56%; oracle ceiling for current candidate set: 83.62%."
        ], 0.85, 5.0, 10.9, 1.0, size=13.2)

    elif kind == "experiment2":
        add_title(slide, "Experiment 2", "SFT/LoRA did not beat controlled repair")
        table(slide, ["Attempt", "Training target", "Observed outcome"], [
            ["Public FRQ LoRA", "751 public FRQ answers", "Overfit behavior; worse on sampled eval"],
            ["OpenR1 LoRA", "External solution traces", "Too much solution-style mismatch"],
            ["Self-distill LoRA", "Base-correct public FRQ, r=8", "Holdout repaired: 110/200 = 55.0%"],
            ["Final decision", "No adapter", "Use base model + repair for FRQ"],
        ], 0.75, 1.45, 11.55, 2.45, col_widths=[2.25, 4.05, 5.25], font_size=10)
        rect(slide, 0.85, 4.55, 10.9, 1.15, fill=RGBColor(250, 247, 242), line=RGBColor(235, 222, 208), radius=True)
        add_text(slide, "Takeaway", 1.15, 4.82, 1.3, 0.25, size=13, color=ORANGE, bold=True)
        add_text(slide, "Fine-tuning is not automatically an accuracy gain. The adapter must preserve reasoning and change only the output contract; otherwise it learns the wrong behavior faster than it improves format.", 2.35, 4.73, 8.8, 0.55, size=13, color=BLACK)

    elif kind == "strengths":
        add_title(slide, "Strength/Weakness/Improvement")
        cols = [
            ("3 Strengths", [
                "Strong baseline separation: MCQ unchanged, FRQ improved separately",
                "Evidence-led repair: every candidate tested with the course judger",
                "Submission validation caught invalid boxes, blanks, and CSV format risk",
            ], GREEN),
            ("3 Weaknesses", [
                "Template repairs can overfit public-style patterns",
                "Long-token thinking runs are slow and memory-sensitive",
                "LoRA attempts lacked enough matched final-answer training data",
            ], RED),
            ("Improvements", [
                "Selective long generation only for uncertain FRQ",
                "Private-safe template coverage with heldout validation",
                "Better final-answer SFT data before trying adapters again",
            ], ACCENT_DARK),
        ]
        for i, (title, bullets, color) in enumerate(cols):
            x = 0.75 + i * 4.1
            add_text(slide, title, x, 1.4, 3.2, 0.3, size=16, bold=True, color=color)
            add_bullets(slide, bullets, x, 1.95, 3.3, 2.75, size=12.5)

    elif kind == "learned":
        add_title(slide, "What have you learned")
        add_bullets(slide, [
            "Small thinking models spend accuracy inside the thinking trace, so max_tokens is part of the model, not just a speed setting.",
            "MCQ and FRQ need different handling: MCQ wants a valid final letter; FRQ can benefit from symbolic cleanup and answer-count repair.",
            "Postprocessing is not a cosmetic layer; for symbolic grading, it can recover correct math that was expressed in the wrong surface form.",
            "A leaderboard improvement is most reliable when each patch has a local validation story and a rollback path."
        ], 0.85, 1.55, 7.2, 2.5, size=15)
        rect(slide, 8.65, 1.55, 3.05, 2.5, fill=RGBColor(244, 247, 248), line=RGBColor(224, 230, 233), radius=True)
        add_text(slide, "Score movement", 9.0, 1.92, 2.35, 0.25, size=12, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(slide, "0.650", 9.0, 2.28, 2.35, 0.35, size=20, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(slide, "to", 9.0, 2.75, 2.35, 0.25, size=11, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(slide, "0.671", 9.0, 3.05, 2.35, 0.42, size=25, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

    elif kind == "future":
        add_title(slide, "Future Work")
        add_bullets(slide, [
            "Run selective second-pass generation on only wrong or low-confidence FRQ instead of regenerating the full set.",
            "Build a private-safe repair library from problem classes rather than memorized public answers.",
            "Train a small adapter only on final-answer formatting after constructing a truly matched external dataset.",
            "Compare Qwen3-4B-Thinking against a non-thinking model under the same wall-clock budget.",
            "Package the final Kaggle notebook so the CSV is produced as an output artifact every time."
        ], 0.85, 1.4, 7.4, 2.9, size=15)
        rect(slide, 8.7, 1.55, 3.15, 2.55, fill=RGBColor(246, 249, 250), line=RGBColor(226, 233, 236), radius=True)
        add_text(slide, "Next target", 9.05, 1.9, 2.45, 0.25, size=12, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(slide, "85%+", 9.05, 2.22, 2.45, 0.48, size=30, color=ACCENT_DARK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, "FRQ accuracy with less public-set leakage", 9.0, 2.86, 2.55, 0.5, size=11, color=GRAY, align=PP_ALIGN.CENTER)


prs.save(PPTX_OUT)


def pdf_text(c, text, x, y, size=18, color=colors.black, bold=False, align="left"):
    c.setFillColor(color)
    c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
    if align == "center":
        c.drawCentredString(x, y, text)
    else:
        c.drawString(x, y, text)


def wrap_lines(text, max_chars):
    words = text.split()
    lines, cur = [], ""
    for word in words:
        nxt = f"{cur} {word}".strip()
        if len(nxt) > max_chars and cur:
            lines.append(cur)
            cur = word
        else:
            cur = nxt
    if cur:
        lines.append(cur)
    return lines


page_w, page_h = W * 72, H * 72
c = canvas.Canvas(str(PDF_OUT), pagesize=(page_w, page_h))
for idx, spec in enumerate(slides, 1):
    c.setFillColor(colors.white)
    c.rect(0, 0, page_w, page_h, stroke=0, fill=1)
    kind = spec["kind"]
    if kind == "title":
        pdf_text(c, spec["title"], page_w/2, page_h*0.56, 28, align="center")
        pdf_text(c, spec["subtitle"], page_w/2, page_h*0.50, 15, colors.HexColor("#5C5C5C"), align="center")
        pdf_text(c, "Qwen3-4B-Thinking | vLLM | FRQ repair | Kaggle submission", page_w/2, page_h*0.44, 10, colors.HexColor("#30404A"), align="center")
    elif kind == "section":
        pdf_text(c, spec["title"], page_w/2, page_h/2, 28, align="center")
    else:
        title = {
            "summary": "Summary", "keywords": "Key Words", "team": "Team Introduction",
            "data": "Data Processing", "model": "Deep Learning Model", "engineering": "Engineering Tricks",
            "experiment1": "Experiment 1", "experiment2": "Experiment 2",
            "strengths": "Strength/Weakness/Improvement", "learned": "What have you learned",
            "future": "Future Work",
        }[kind]
        pdf_text(c, title, 45, page_h-55, 17)
        if kind == "summary":
            y = page_h-105
            for line in spec["body"]:
                for wrapped in wrap_lines(line, 90):
                    pdf_text(c, "• " + wrapped, 60, y, 12)
                    y -= 19
            pdf_text(c, "Final public score: 0.671    Public FRQ repair: 82.56%", 60, 120, 15, colors.HexColor("#2C8452"), bold=True)
        elif kind == "keywords":
            pdf_text(c, "Qwen3 Thinking Model    vLLM    FRQ Repair    Token Budget    Kaggle", page_w/2, page_h/2, 23, colors.HexColor("#708792"), align="center")
        else:
            # Compact fallback PDF rendering; PPTX is the primary editable output.
            y = page_h-105
            text_map = {
                "team": ["Keith Gong - Department of Data Science, UC San Diego", "Owned inference pipeline, prompt experiments, repair functions, and Kaggle submission assembly."],
                "data": ["Public: 1126 rows (375 MCQ, 751 FRQ). Private: 943 rows (300 MCQ, 643 FRQ).", "Pipeline: prompt -> vLLM -> boxed extraction -> FRQ repair -> CSV."],
                "model": ["Qwen3-4B-Thinking-2507 with vLLM and BitsAndBytes quantization.", "MCQ kept baseline unchanged; FRQ used base generation plus deterministic repair."],
                "engineering": ["Remote nohup jobs, log files, PID checks, nvidia-smi monitoring.", "JSONL outputs and validation prevented silent submission failures."],
                "experiment1": ["Milestone baseline: 71.1% MCQ, 54.8% FRQ.", "Full public FRQ: 56.6% raw -> 82.6% final repair."],
                "experiment2": ["Public FRQ LoRA, OpenR1 LoRA, and self-distill LoRA failed to beat controlled repair.", "Final decision: no adapter in submitted pipeline."],
                "strengths": ["Strengths: separated MCQ/FRQ, evidence-led repair, submission validation.", "Weaknesses: public-style repair can overfit; token budget remains slow."],
                "learned": ["Token budget is part of model behavior.", "Postprocessing is a measurable accuracy layer for symbolic grading."],
                "future": ["Selective second-pass generation, private-safe repair library, matched final-answer SFT data.", "Target: 85%+ FRQ with less public-set leakage."],
            }
            for line in text_map.get(kind, []):
                for wrapped in wrap_lines(line, 95):
                    pdf_text(c, "• " + wrapped, 60, y, 13)
                    y -= 22
    pdf_text(c, str(idx), page_w-35, 24, 8, colors.HexColor("#9A9A9A"))
    c.showPage()
c.save()

print(PPTX_OUT)
print(PDF_OUT)
